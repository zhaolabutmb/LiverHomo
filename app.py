import base64
import gc
import io
import json
import os
import re
import smtplib
import threading
from collections import deque
from email.message import EmailMessage
from typing import Any
from urllib import error as urllib_error
from urllib import request as urllib_request

import anndata as ad
import matplotlib
import networkx as nx
import numpy as np
import pandas as pd
import seaborn as sns
import tiktoken
from flask import Flask, jsonify, render_template, request, send_file, url_for
from openpyxl import load_workbook
from assistant_data import DegIndex, NameResolver, load_composition, short_disease_name
from env_config import load_env
from expression_store import ExpressionStore
from umap_config import UMAP_METADATA_COLUMNS, UMAP_STATIC_IMAGE_FILES

matplotlib.use("Agg")
import matplotlib.pyplot as plt

load_env()

METADATA_PATH = os.environ["METADATA_PATH"]
FIGS_ROOT = os.environ["FIGS_ROOT"]
GRAPH_ROOT = os.environ["GRAPH_ROOT"]
DEG_ROOT = os.environ["DEG_ROOT"]
VOLCANO_ROOT = os.environ["VOLCANO_ROOT"]
EGO_ROOT = os.environ["EGO_ROOT"]
CELL_INTERACTION_ROOT = os.environ["CELL_INTERACTION_ROOT"]
GENE_EXPRESSION_DATA_PATH = os.environ["GENE_EXPRESSION_DATA_PATH"]
GENE_VECTOR_CACHE_DIR = os.environ["GENE_VECTOR_CACHE_DIR"]
OPENAI_API_KEY = os.environ["OPENAI_API_KEY"]
CONTACT_RECIPIENT = "zhaoresearchlab@gmail.com"
SMTP_HOST = os.environ.get("SMTP_HOST", "")
SMTP_PORT = int(os.environ.get("SMTP_PORT", "587"))
SMTP_USERNAME = os.environ.get("SMTP_USERNAME", "")
SMTP_PASSWORD = os.environ.get("SMTP_PASSWORD", "")
SMTP_FROM_EMAIL = os.environ.get("SMTP_FROM_EMAIL", SMTP_USERNAME)
GRAPH_LAYOUT_EDGE_LIMIT = None
GRAPH_LAYOUTS = {"spring", "kamada_kawai", "random"}
GRAPH_DATA_CACHE: dict[str, dict[str, Any]] = {}
GRAPH_DATA_CACHE_LOCK = threading.Lock()
GRAPH_DATA_LOADING: set[str] = set()
GRAPH_DATA_ERRORS: dict[str, str] = {}

app = Flask(__name__)
app.json.sort_keys = False
FILTER_CATALOG_PATH = os.path.join(app.root_path, "static", "filter_catalog.json")

RELOAD = os.environ.get("RELOAD", "false").lower() in {"1", "true", "yes", "on"}
app.config["TEMPLATES_AUTO_RELOAD"] = RELOAD
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0 if RELOAD else None
_adata = None
_gene_expression_adata = None
BACKGROUND_DATA = {"summary": None, "graph_summary": None, "loading": False, "error": None}
BACKGROUND_DATA_LOCK = threading.Lock()
GENE_VECTOR_CACHE: dict[int, np.ndarray] = {}
GENE_VECTOR_CACHE_LOCK = threading.Lock()
GENE_VECTOR_CACHE_SIZE = 16
OBS_COLUMN_CACHE: dict[str, np.ndarray] = {}
OBS_COLUMN_CACHE_LOCK = threading.Lock()
GENE_PLOT_CACHE: dict[tuple[str, str, str, str], str] = {}
GENE_PLOT_CACHE_LOCK = threading.Lock()
GENE_PLOT_CACHE_SIZE = 32
GENE_PLOT_MAX_PER_GROUP = 1_500
GENE_PLOT_MAX_POINTS = 60_000


def load_adata():
    global _adata
    if _adata is None:
        _adata = ad.read_h5ad(METADATA_PATH, backed="r")
    return _adata


def load_gene_expression_adata():
    global _gene_expression_adata
    if _gene_expression_adata is None:
        if not os.path.isfile(GENE_EXPRESSION_DATA_PATH):
            raise FileNotFoundError(
                f"Condensed gene-expression AnnData not found at {GENE_EXPRESSION_DATA_PATH}. "
                "Generate it with scripts/generate_gene_expression_adata.py."
            )
        _gene_expression_adata = ExpressionStore(GENE_EXPRESSION_DATA_PATH)
    return _gene_expression_adata


def load_filter_catalog() -> dict[str, Any]:
    with open(FILTER_CATALOG_PATH, encoding="utf-8") as handle:
        return json.load(handle)


@app.context_processor
def inject_filter_catalog():
    catalog = load_filter_catalog()
    catalog["umap_columns"] = get_static_umap_metadata_columns()
    return {"filter_catalog": catalog}


def unload_adata():
    global _adata, _gene_expression_adata
    if _adata is not None:
        _adata.file.close()
        _adata = None
    if _gene_expression_adata is not None:
        _gene_expression_adata.close()
        _gene_expression_adata = None
    with BACKGROUND_DATA_LOCK:
        BACKGROUND_DATA["summary"] = None
    with GENE_VECTOR_CACHE_LOCK:
        GENE_VECTOR_CACHE.clear()
    with OBS_COLUMN_CACHE_LOCK:
        OBS_COLUMN_CACHE.clear()
    with GENE_PLOT_CACHE_LOCK:
        GENE_PLOT_CACHE.clear()
    gc.collect()


def unload_graph_data():
    with BACKGROUND_DATA_LOCK:
        BACKGROUND_DATA["graph_summary"] = None
    gc.collect()


def list_graph_files():
    if not os.path.isdir(GRAPH_ROOT):
        return []
    files = []
    for entry in sorted(os.listdir(GRAPH_ROOT)):
        path = os.path.join(GRAPH_ROOT, entry)
        if os.path.isfile(path) and entry.endswith(".xlsx"):
            files.append(path)
    return files


def parse_graph_label(path: str) -> dict[str, str]:
    name = os.path.basename(path)
    stem = name.replace("_querylist1.xlsx", "")
    if stem.endswith("_querylist1"):
        stem = stem[: -len("_querylist1")]
    parts = stem.split("_", 1)
    disease = parts[0] if parts else "Unknown"
    cell_type = parts[1] if len(parts) > 1 else "Unknown"
    return {"disease": disease, "cell_type": cell_type, "file": name, "path": path}


def load_graph_data(path: str, max_edges: int | None = None) -> dict[str, Any]:
    cache_key = f"{path}:{max_edges or 0}"
    with GRAPH_DATA_CACHE_LOCK:
        cached = GRAPH_DATA_CACHE.get(cache_key)
        if cached is not None:
            return cached

    if path.lower().endswith(".txt"):
        try:
            frame = pd.read_csv(path, sep="\t", header=None, names=["geneID1", "geneID2"], engine="python")
        except pd.errors.EmptyDataError:
            return {"nodes": [], "edges": []}
        frame = frame.dropna(subset=["geneID1", "geneID2"]).copy()
        frame["geneID1"] = frame["geneID1"].astype(str).str.strip()
        frame["geneID2"] = frame["geneID2"].astype(str).str.strip()
        frame = frame[frame["geneID1"] != ""][frame["geneID2"] != ""]
        if frame.empty:
            return {"nodes": [], "edges": []}
        combined = frame.drop_duplicates(subset=["geneID1", "geneID2"])
        nodes = sorted(set(combined["geneID1"].tolist()) | set(combined["geneID2"].tolist()))
        edges = [{"source": str(row["geneID1"]), "target": str(row["geneID2"]), "score": 1.0} for _, row in combined.iterrows()]
        payload = {"nodes": nodes, "edges": edges}
        with GRAPH_DATA_CACHE_LOCK:
            GRAPH_DATA_CACHE[cache_key] = payload
        return payload

    edges_by_key = {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
            columns = {name: headers.index(name) for name in ("geneID1", "geneID2", "Score") if name in headers}
            if len(columns) != 3:
                continue
            for row in rows:
                source = row[columns["geneID1"]]
                target = row[columns["geneID2"]]
                score = row[columns["Score"]]
                if source is None or target is None or score is None:
                    continue
                key = (str(source), str(target), str(score))
                edges_by_key[key] = {"source": str(source), "target": str(target), "score": float(score)}
    finally:
        workbook.close()

    edges = list(edges_by_key.values())
    if max_edges is not None and len(edges) > max_edges:
        edges.sort(key=lambda edge: edge["score"], reverse=True)
        edges = edges[:max_edges]
    nodes = sorted({edge["source"] for edge in edges} | {edge["target"] for edge in edges})
    payload = {"nodes": nodes, "edges": edges}
    with GRAPH_DATA_CACHE_LOCK:
        GRAPH_DATA_CACHE[cache_key] = payload
    return payload


def get_graph_catalog() -> list[dict[str, Any]]:
    return [parse_graph_label(path) for path in list_graph_files()]


def start_graph_data_load(path: str, max_edges: int | None = None):
    cache_key = f"{path}:{max_edges or 0}"
    with GRAPH_DATA_CACHE_LOCK:
        if cache_key in GRAPH_DATA_CACHE or cache_key in GRAPH_DATA_LOADING:
            return
        GRAPH_DATA_LOADING.add(cache_key)

    def load():
        try:
            load_graph_data(path, max_edges=max_edges)
        except Exception as exc:  # pragma: no cover - depends on external graph files
            with GRAPH_DATA_CACHE_LOCK:
                GRAPH_DATA_ERRORS[cache_key] = str(exc)
        finally:
            with GRAPH_DATA_CACHE_LOCK:
                GRAPH_DATA_LOADING.discard(cache_key)

    threading.Thread(target=load, name="livertraj-graph-load", daemon=True).start()


def graph_data_status(path: str, max_edges: int | None = None) -> str:
    cache_key = f"{path}:{max_edges or 0}"
    with GRAPH_DATA_CACHE_LOCK:
        if cache_key in GRAPH_DATA_ERRORS:
            return "error"
        if cache_key in GRAPH_DATA_CACHE:
            return "ready"
        if cache_key in GRAPH_DATA_LOADING:
            return "loading"
    return "missing"


def get_graph_summary() -> dict[str, Any]:
    catalog = get_graph_catalog()
    diseases = sorted({item["disease"] for item in catalog})
    cell_types = sorted({item["cell_type"] for item in catalog})
    return {"files": len(catalog), "diseases": diseases, "cell_types": cell_types}


def get_ego_catalog() -> list[dict[str, str]]:
    if not os.path.isdir(EGO_ROOT):
        return []

    catalog = []
    for disease in sorted(os.listdir(EGO_ROOT)):
        disease_dir = os.path.join(EGO_ROOT, disease)
        if not os.path.isdir(disease_dir):
            continue
        for filename in sorted(os.listdir(disease_dir)):
            if not filename.endswith("_ego_ppi.png"):
                continue
            prefix = f"{disease}_"
            if not filename.startswith(prefix):
                continue
            cell_type = filename[len(prefix):-len("_ego_ppi.png")]
            catalog.append({"disease": disease, "cell_type": cell_type, "filename": filename})
    return catalog


def get_ego_image_path(disease: str, cell_type: str) -> str | None:
    allowed = {
        (item["disease"], item["cell_type"]): item["filename"]
        for item in get_ego_catalog()
    }
    filename = allowed.get((disease, cell_type))
    if filename is None:
        return None
    return os.path.join(EGO_ROOT, disease, filename)


CELL_INTERACTION_KIND_SUFFIXES = {"heatmap": "_cci_heatmap.png", "chord": "_cci_chord.png"}
CELL_INTERACTION_COUNTS_SUFFIX = "_interaction_counts.csv"


def get_cell_interaction_catalog(kind: str = "heatmap") -> list[dict[str, str]]:
    if not os.path.isdir(CELL_INTERACTION_ROOT):
        return []

    suffix = CELL_INTERACTION_KIND_SUFFIXES.get(kind, CELL_INTERACTION_KIND_SUFFIXES["heatmap"])
    catalog = []
    for filename in sorted(os.listdir(CELL_INTERACTION_ROOT)):
        if not filename.endswith(suffix):
            continue
        catalog.append({"disease": filename[: -len(suffix)], "filename": filename})
    return catalog


def get_cell_interaction_path(disease: str, kind: str = "heatmap") -> str | None:
    allowed = {item["disease"]: item["filename"] for item in get_cell_interaction_catalog(kind)}
    filename = allowed.get(disease)
    if filename is None:
        return None
    return os.path.join(CELL_INTERACTION_ROOT, filename)


def load_cell_interaction_matrix(disease: str) -> pd.DataFrame | None:
    if not os.path.isdir(CELL_INTERACTION_ROOT):
        return None
    safe_name = os.path.basename(disease)
    path = os.path.join(CELL_INTERACTION_ROOT, f"{safe_name}{CELL_INTERACTION_COUNTS_SUFFIX}")
    if not os.path.isfile(path):
        return None
    matrix = pd.read_csv(path, index_col=0)
    # The CSVs write "Circulating_NK/NKT"; the rest of the atlas uses "Circulating_NK_NKT".
    matrix.index = matrix.index.str.replace("/", "_", regex=False)
    matrix.columns = matrix.columns.str.replace("/", "_", regex=False)
    return matrix


def get_cell_interaction_cell_types() -> list[str]:
    for item in get_cell_interaction_catalog("heatmap"):
        matrix = load_cell_interaction_matrix(item["disease"])
        if matrix is not None:
            return sorted(str(name) for name in matrix.columns)
    return []


def get_cell_interaction_pair_counts(cell_type_a: str, cell_type_b: str) -> dict[str, Any]:
    cell_type_a = (cell_type_a or "").strip()
    cell_type_b = (cell_type_b or "").strip()
    if not cell_type_a or not cell_type_b:
        raise ValueError("Two cell types are required.")

    rows = []
    for item in get_cell_interaction_catalog("heatmap"):
        disease = item["disease"]
        matrix = load_cell_interaction_matrix(disease)
        if matrix is None or cell_type_a not in matrix.index or cell_type_b not in matrix.columns:
            continue
        count = matrix.loc[cell_type_a, cell_type_b]
        rows.append({"disease": disease, "count": float(count) if pd.notna(count) else 0.0})
    rows.sort(key=lambda row: row["disease"])

    return {
        "cell_type_a": cell_type_a,
        "cell_type_b": cell_type_b,
        "rows": rows,
        "image_base64": _get_cell_interaction_histogram_image(cell_type_a, cell_type_b, rows),
    }


def _get_cell_interaction_histogram_image(cell_type_a: str, cell_type_b: str, rows: list[dict[str, Any]]) -> str:
    sns.set_theme(style="whitegrid", font_scale=1.2)
    if not rows:
        fig, ax = plt.subplots(figsize=(8, 4))
        ax.text(0.5, 0.5, "No interaction data for the selected cell types", ha="center", va="center")
        ax.axis("off")
        return _encode_plot_png(fig)

    frame = pd.DataFrame(rows)
    fig, ax = plt.subplots(figsize=(max(6, len(frame) * 0.7), 5))
    sns.barplot(data=frame, x="disease", y="count", hue="disease", legend=False, palette="Set2", ax=ax)
    ax.set_xlabel("Disease")
    ax.set_ylabel("Interaction count")
    ax.set_title(f"{cell_type_a} \u2194 {cell_type_b} interactions by disease")
    ax.tick_params(axis="x", rotation=45)
    for label in ax.get_xticklabels():
        label.set_ha("right")
    return _encode_plot_png(fig)


def get_volcano_catalog() -> list[dict[str, str]]:
    if not os.path.isdir(VOLCANO_ROOT):
        return []

    preferred: dict[str, str] = {}
    fallback: dict[str, str] = {}
    for filename in sorted(os.listdir(VOLCANO_ROOT)):
        if not filename.lower().endswith(".png"):
            continue
        if filename.startswith("DEG_volcano_"):
            disease = filename[len("DEG_volcano_"):-len(".png")]
            preferred[disease] = filename
        elif filename.endswith("_volcano.png"):
            disease = filename[:-len("_volcano.png")]
            fallback[disease] = filename

    diseases = sorted(set(preferred) | set(fallback))
    return [{"disease": disease, "filename": preferred.get(disease, fallback[disease])} for disease in diseases]


def get_volcano_path(disease: str) -> str | None:
    allowed = {item["disease"]: item["filename"] for item in get_volcano_catalog()}
    filename = allowed.get(disease)
    if filename is None:
        return None
    return os.path.join(VOLCANO_ROOT, filename)


def parse_gene_list(raw_gene_list: str | None) -> list[str]:
    if raw_gene_list is None:
        return []
    genes = []
    for value in re.split(r"[\n,;\t]+", raw_gene_list):
        gene = value.strip()
        if gene:
            genes.append(gene)
    return genes


def filter_graph_by_genes(graph_payload: dict[str, Any], gene_list: list[str]) -> dict[str, Any]:
    if not gene_list:
        return graph_payload

    normalized_genes = {str(gene).strip().lower() for gene in gene_list if str(gene).strip()}
    if not normalized_genes:
        return graph_payload

    graph = nx.Graph()
    for node in graph_payload.get("nodes", []):
        node_name = str(node).strip()
        if node_name:
            graph.add_node(node_name)

    for edge in graph_payload.get("edges", []):
        source = str(edge.get("source", "")).strip()
        target = str(edge.get("target", "")).strip()
        if source and target and source != target:
            graph.add_edge(source, target)

    seed_nodes = [node for node in graph.nodes() if node.lower() in normalized_genes]
    if not seed_nodes:
        return {"nodes": [], "edges": []}

    subgraph_nodes = set(seed_nodes)
    for node in seed_nodes:
        subgraph_nodes.update(graph.neighbors(node))

    seed_node_set = set(seed_nodes)
    connected_edges = []
    for edge in graph_payload.get("edges", []):
        source = str(edge.get("source", "")).strip()
        target = str(edge.get("target", "")).strip()
        if (
            source in subgraph_nodes
            and target in subgraph_nodes
            and (source in seed_node_set or target in seed_node_set)
        ):
            connected_edges.append(edge)

    return {"nodes": sorted(subgraph_nodes), "edges": connected_edges}


def get_graph_payload(
    disease: str | None = None,
    cell_type: str | None = None,
    limit: int = 200,
    gene_list: str | list[str] | None = None,
    max_edges: int | None = None,
) -> dict[str, Any]:
    genes = parse_gene_list(gene_list if isinstance(gene_list, str) else ",".join(gene_list or []))
    selected_diseases = _parse_multi_select_values(disease)
    selected_cell_types = _parse_multi_select_values(cell_type)

    catalog = get_graph_catalog()
    filtered = catalog
    if selected_diseases:
        disease_filters = {value.lower() for value in selected_diseases}
        filtered = [item for item in filtered if any(item["disease"].lower() == value for value in disease_filters)]
    if selected_cell_types:
        cell_type_filters = {value.lower() for value in selected_cell_types}
        filtered = [item for item in filtered if any(item["cell_type"].lower() == value for value in cell_type_filters)]
    selected = filtered[:limit]
    graphs = []
    for item in selected:
        payload = load_graph_data(item["path"], max_edges=max_edges)
        if genes:
            payload = filter_graph_by_genes(payload, genes)
        graphs.append({"file": item["file"], "disease": item["disease"], "cell_type": item["cell_type"], **payload})
    return {"catalog": selected, "graphs": graphs}


def build_graph_network(graph_payload: dict[str, Any], layout: str = "random") -> dict[str, Any]:
    if layout not in GRAPH_LAYOUTS:
        raise ValueError(f"Unknown graph layout: {layout}")

    nodes = graph_payload.get("nodes", []) or []
    edges = graph_payload.get("edges", []) or []
    if GRAPH_LAYOUT_EDGE_LIMIT is not None and len(edges) > GRAPH_LAYOUT_EDGE_LIMIT:
        edges = sorted(edges, key=lambda edge: float(edge.get("score", 0.0) or 0.0), reverse=True)[:GRAPH_LAYOUT_EDGE_LIMIT]
    if not nodes and not edges:
        return {"nodes": [], "edges": [], "positions": {}}

    graph = nx.Graph()
    for edge in edges:
        source = str(edge.get("source", ""))
        target = str(edge.get("target", ""))
        if not source or not target or source == target:
            continue
        graph.add_edge(source, target, weight=float(edge.get("score", 0.0) or 0.0))
    if graph.number_of_nodes() == 0:
        limit = GRAPH_LAYOUT_EDGE_LIMIT if GRAPH_LAYOUT_EDGE_LIMIT is not None else len(nodes)
        for node_name in nodes[:limit]:
            graph.add_node(str(node_name))

    if graph.number_of_nodes() == 0:
        return {"nodes": [], "edges": [], "positions": {}}

    if layout == "spring":
        positions = nx.spring_layout(graph, seed=0)
    elif layout == "kamada_kawai":
        positions = nx.kamada_kawai_layout(graph)
    else:
        positions = nx.random_layout(graph, seed=0)
    x_values = [float(position[0]) for position in positions.values()]
    y_values = [float(position[1]) for position in positions.values()]
    min_x, max_x = min(x_values), max(x_values)
    min_y, max_y = min(y_values), max(y_values)
    x_span = max(max_x - min_x, 1e-9)
    y_span = max(max_y - min_y, 1e-9)
    layout_nodes = []
    for node_name in sorted(graph.nodes()):
        x, y = positions[node_name]
        layout_nodes.append({
            "id": str(node_name),
            "x": float((x - min_x) / x_span),
            "y": float((y - min_y) / y_span),
            "degree": int(graph.degree(node_name)),
        })

    layout_edges = []
    for source, target, data in graph.edges(data=True):
        layout_edges.append({
            "source": str(source),
            "target": str(target),
            "score": float(data.get("weight", 0.0) or 0.0),
        })

    return {"nodes": layout_nodes, "edges": layout_edges, "positions": {node["id"]: {"x": node["x"], "y": node["y"]} for node in layout_nodes}}


def get_graph_network_payload(
    disease: str | None = None,
    cell_type: str | None = None,
    limit: int = 200,
    gene_list: str | list[str] | None = None,
    layout: str = "random",
) -> dict[str, Any]:
    if layout not in GRAPH_LAYOUTS:
        raise ValueError(f"Unknown graph layout: {layout}")

    selected_diseases = _parse_multi_select_values(disease)
    selected_cell_types = _parse_multi_select_values(cell_type)
    catalog = get_graph_catalog()
    if selected_diseases:
        allowed_diseases = {value.lower() for value in selected_diseases}
        catalog = [item for item in catalog if item["disease"].lower() in allowed_diseases]
    if selected_cell_types:
        allowed_cell_types = {value.lower() for value in selected_cell_types}
        catalog = [item for item in catalog if item["cell_type"].lower() in allowed_cell_types]
    selected = catalog[:1]
    if selected:
        item = selected[0]
        status = graph_data_status(item["path"], max_edges=GRAPH_LAYOUT_EDGE_LIMIT)
        graph_info = {"file": item["file"], "disease": item["disease"], "cell_type": item["cell_type"]}
        if status == "missing":
            start_graph_data_load(item["path"], max_edges=GRAPH_LAYOUT_EDGE_LIMIT)
            return {"loading": True, "graph": graph_info}
        if status == "loading":
            return {"loading": True, "graph": graph_info}
        if status == "error":
            cache_key = f"{item['path']}:{GRAPH_LAYOUT_EDGE_LIMIT}"
            with GRAPH_DATA_CACHE_LOCK:
                error = GRAPH_DATA_ERRORS.get(cache_key, "Graph data could not be loaded.")
            return {"error": error, "graph": graph_info}

    payload = get_graph_payload(
        disease=disease,
        cell_type=cell_type,
        limit=1,
        gene_list=gene_list,
        max_edges=GRAPH_LAYOUT_EDGE_LIMIT,
    )
    if not payload.get("graphs"):
        return {"catalog": payload.get("catalog", []), "graph": None, "nodes": [], "edges": []}

    primary_graph = payload["graphs"][0]
    network = build_graph_network(
        {"nodes": primary_graph.get("nodes", []), "edges": primary_graph.get("edges", [])},
        layout=layout,
    )
    return {
        "catalog": payload.get("catalog", []),
        "graph": {"file": primary_graph.get("file"), "disease": primary_graph.get("disease"), "cell_type": primary_graph.get("cell_type")},
        "nodes": network["nodes"],
        "edges": network["edges"],
        "positions": network["positions"],
    }


def get_deg_catalog() -> list[dict[str, str]]:
    if not os.path.isdir(DEG_ROOT):
        return []
    catalog = []
    for disease_dir in sorted(os.listdir(DEG_ROOT)):
        disease_path = os.path.join(DEG_ROOT, disease_dir)
        if not os.path.isdir(disease_path):
            continue
        for filename in sorted(os.listdir(disease_path)):
            if not filename.lower().endswith(".csv") or not filename.startswith("DEG_"):
                continue
            path = os.path.join(disease_path, filename)
            stem = filename[:-4]
            cell_type = stem[len("DEG_" + disease_dir + "_"):] if stem.startswith("DEG_" + disease_dir + "_") else stem.replace("DEG_" + disease_dir + "_", "")
            catalog.append({
                "disease": disease_dir,
                "cell_type": cell_type,
                "file": filename,
                "path": path,
            })
    return catalog


def get_deg_summary() -> dict[str, Any]:
    catalog = get_deg_catalog()
    diseases = sorted({item["disease"] for item in catalog})
    cell_types = sorted({item["cell_type"] for item in catalog})
    return {"diseases": diseases, "cell_types": cell_types, "files": len(catalog)}

DEG_DISPLAY_COLUMNS = [
    "gene",
    "cell_type",
    "protein",
    "sig",
    "de_pval",
    "de_coef",
    "de_se",
    "dv_pval",
    "dv_coef",
    "dv_se",
    "baseMean",
]
DEG_ALL_CELL_TYPES = "__all__"


def get_deg_table(
    disease: str | None = None,
    cell_type: str | None = None,
    regulation: str | None = None,
    limit: int = 0,
    sort_by: str | None = None,
    sort_order: str = "desc",
) -> dict[str, Any]:
    catalog = get_deg_catalog()
    selected_diseases = _parse_multi_select_values(disease)
    selected_cell_types = _parse_multi_select_values(cell_type)
    all_cell_types = DEG_ALL_CELL_TYPES in {value.lower() for value in selected_cell_types}
    if all_cell_types:
        selected_cell_types = []

    filtered = catalog
    if selected_diseases:
        allowed_diseases = {value.lower() for value in selected_diseases}
        filtered = [item for item in filtered if item["disease"].lower() in allowed_diseases]
    if selected_cell_types:
        allowed_cell_types = {value.lower() for value in selected_cell_types}
        filtered = [item for item in filtered if item["cell_type"].lower() in allowed_cell_types]

    if selected_cell_types:
        sources = [item["path"] for item in filtered]
    else:
        # All cell types (or no selection): use the disease-level aggregate files.
        sources = []
        for disease_name in sorted({item["disease"] for item in filtered}):
            aggregate = os.path.join(DEG_ROOT, f"DEG_{disease_name}.csv")
            if os.path.isfile(aggregate):
                sources.append(aggregate)

    frames = []
    for path in sources:
        try:
            frame = pd.read_csv(path, sep=None, engine="python")
        except Exception:
            continue
        if frame.empty:
            continue
        frame = frame.copy()
        if regulation:
            if "sig" not in frame.columns:
                continue
            allowed_regulation = {value.strip().lower() for value in regulation.split(",") if value.strip()}
            frame = frame[frame["sig"].astype(str).str.strip().str.lower().isin(allowed_regulation)]
            if frame.empty:
                continue
        if all_cell_types:
            frame = frame.rename(columns={"cell_types": "cell_type"})
        else:
            frame = frame.drop(columns=["cell_type", "cell_types"], errors="ignore")
        frames.append(frame)

    combined = pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()
    sort_value = (sort_by or "").strip().lower()
    sort_column = sort_value if sort_value in combined.columns else None
    ascending = str(sort_order).strip().lower() == "asc"
    if sort_column and not combined.empty and sort_column in combined.columns:
        sort_key = pd.to_numeric(combined[sort_column], errors="coerce")
        if sort_key.notna().any():
            order = sort_key.sort_values(ascending=ascending, na_position="last").index
        else:
            order = combined[sort_column].astype(str).sort_values(ascending=ascending, na_position="last").index
        combined = combined.loc[order]
    if limit and not combined.empty:
        combined = combined.head(max(1, int(limit)))
    if not combined.empty:
        combined = combined.reindex(columns=[column for column in DEG_DISPLAY_COLUMNS if column in combined.columns])
    return {
        "count": int(len(combined)),
        "rows": combined.to_dict(orient="records"),
        "diseases": sorted({item["disease"] for item in filtered}),
        "cell_types": sorted({item["cell_type"] for item in filtered}),
    }


def get_obs_columns():
    adata = load_adata()
    return list(adata.obs.columns)


def normalize_series(series: Any) -> pd.Series:
    return pd.Series(series).astype("string").fillna("Unknown").astype(str)


def resolve_cell_type_column(preferred: str | None = None) -> str | None:
    adata = load_adata()
    if preferred and preferred in adata.obs.columns:
        return preferred
    if "cell_type2" in adata.obs.columns:
        return "cell_type2"
    return None


def get_cell_type_options(column: str | None = None) -> list[str]:
    adata = load_adata()
    column = column or resolve_cell_type_column()
    if not column or column not in adata.obs.columns:
        return []
    values = normalize_series(adata.obs[column]).value_counts().head(20).index.tolist()
    return values


def get_disease_options() -> list[str]:
    adata = load_adata()
    if "diagnosis" not in adata.obs.columns:
        return []
    return normalize_series(adata.obs["diagnosis"]).value_counts().head(20).index.tolist()


def get_umap_metadata_columns() -> list[str]:
    adata = load_adata()
    available = set(adata.obs.columns)
    return [name for name in UMAP_METADATA_COLUMNS if name in available]


def get_static_umap_metadata_columns() -> list[str]:
    return [
        name for name in UMAP_STATIC_IMAGE_FILES
        if os.path.isfile(get_static_umap_image_path(name))
    ]


def get_static_umap_image_path(column: str) -> str | None:
    filename = UMAP_STATIC_IMAGE_FILES.get(column)
    if filename is None:
        return None
    return os.path.join(FIGS_ROOT, filename)


def summarize_dataset():
    adata = load_adata()
    filter_catalog = load_filter_catalog()
    summary = {
        "cells": int(adata.n_obs),
        "genes": int(load_gene_expression_adata().n_vars),
        "obs_columns": len(adata.obs.columns),
        "sample_columns": list(adata.obs.columns[:12]),
        "dataset_counts": {},
        "type_counts": {},
        "cell_type_column": resolve_cell_type_column(),
        "cell_type_options": filter_catalog["metadata"]["cell_types"],
        "disease_column": "diagnosis" if "diagnosis" in adata.obs.columns else None,
        "disease_options": filter_catalog["metadata"]["diseases"],
        "umap_columns": get_static_umap_metadata_columns(),
    }

    if "dataset" in adata.obs.columns:
        summary["dataset_counts"] = normalize_series(adata.obs["dataset"]).value_counts().head(8).to_dict()
    if "Type" in adata.obs.columns:
        summary["type_counts"] = normalize_series(adata.obs["Type"]).value_counts().head(8).to_dict()
    if "cell_type2" in adata.obs.columns:
        summary["cell_type_counts"] = normalize_series(adata.obs["cell_type2"]).value_counts().head(8).to_dict()

    return summary


def load_all_data_in_background():
    global BACKGROUND_DATA
    with BACKGROUND_DATA_LOCK:
        if BACKGROUND_DATA["loading"]:
            return
        BACKGROUND_DATA["loading"] = True
        BACKGROUND_DATA["error"] = None

    try:
        summary = summarize_dataset()
        with BACKGROUND_DATA_LOCK:
            BACKGROUND_DATA["summary"] = summary
            BACKGROUND_DATA["error"] = None
    except Exception as exc:  # pragma: no cover - background loader failure is non-fatal
        with BACKGROUND_DATA_LOCK:
            BACKGROUND_DATA["error"] = str(exc)
    finally:
        with BACKGROUND_DATA_LOCK:
            BACKGROUND_DATA["loading"] = False


def start_background_data_load():
    worker = threading.Thread(target=load_all_data_in_background, name="livertraj-background-load", daemon=True)
    worker.start()
    return worker


def search_rows(
    column: str | None,
    value: str | None,
    limit: int = 100,
    cell_type: str | None = None,
    disease: str | None = None,
    sort_by: str | None = None,
    sort_order: str = "asc",
):
    adata = load_adata()
    obs = adata.obs
    subset = obs

    if column and value:
        query = str(value).strip()
        if column in obs.columns:
            subset = obs.loc[normalize_series(obs[column]).str.contains(query, case=False, na=False)]
        else:
            subset = obs.iloc[0:0]

    effective_column = "cell_type2" if "cell_type2" in obs.columns else None
    selected_cell_types = _parse_multi_select_values(cell_type)
    if selected_cell_types and effective_column:
        filters = pd.Series(False, index=subset.index)
        for category in selected_cell_types:
            filters |= normalize_series(subset[effective_column]).str.contains(str(category).strip(), case=False, na=False)
        subset = subset.loc[filters]

    selected_diseases = _parse_multi_select_values(disease)
    if selected_diseases and "diagnosis" in obs.columns:
        filters = pd.Series(False, index=subset.index)
        for category in selected_diseases:
            filters |= normalize_series(subset["diagnosis"]).str.contains(str(category).strip(), case=False, na=False)
        subset = subset.loc[filters]

    sort_column = (sort_by or "").strip()
    if sort_column and sort_column in subset.columns:
        ascending = str(sort_order).strip().lower() != "desc"
        sort_key = pd.to_numeric(subset[sort_column], errors="coerce")
        if not sort_key.notna().any():
            sort_key = normalize_series(subset[sort_column])
        subset = subset.loc[sort_key.sort_values(ascending=ascending, na_position="last").index]

    subset = subset.head(limit)
    display_columns = [name for name in get_umap_metadata_columns() if name in subset.columns]
    leading = [name for name in ("diagnosis", "cell_type2") if name in display_columns]
    display_columns = leading + [name for name in display_columns if name not in leading]
    subset = subset[display_columns].reset_index()
    index_name = subset.columns[0]
    subset = subset[[*display_columns, index_name]]
    return subset.to_dict(orient="records")


def _gene_lookup(adata: ExpressionStore) -> dict[str, tuple[int, str, str]]:
    # Upper-cased gene ID or symbol -> (index, display name, matrix). Main
    # (lognorm) genes win over raw-only genes; IDs win over gene_name symbols.
    lookup = getattr(adata, "_gene_lookup", None)
    if lookup is not None:
        return lookup
    lookup = {}
    sources = []
    if adata.raw is not None:
        raw_names = pd.Index(adata.raw.var_names).astype(str)
        if "gene_name" in adata.raw.var.columns:
            sources.append(("raw", raw_names, normalize_series(adata.raw.var["gene_name"])))
        sources.append(("raw", raw_names, raw_names))
    var_names = pd.Index(adata.var_names).astype(str)
    if "gene_name" in adata.var.columns:
        sources.append(("main", var_names, normalize_series(adata.var["gene_name"])))
    sources.append(("main", var_names, var_names))
    for source, names, keys in sources:  # later sources overwrite earlier ones
        for idx, key in enumerate(keys):
            lookup[str(key).upper()] = (idx, str(names[idx]), source)
    adata._gene_lookup = lookup
    return lookup


def _resolve_gene_index(adata: ExpressionStore, gene_query: str) -> tuple[int, str, str]:
    gene_query = (gene_query or "").strip()
    if not gene_query:
        raise ValueError("A gene name is required")
    lookup = _gene_lookup(adata)
    match = lookup.get(gene_query.upper())
    if match is not None:
        return match
    query = gene_query.upper()
    suggestions = sorted({lookup[key][1] for key in lookup if key.startswith(query)})[:8]
    hint = f" Did you mean: {', '.join(suggestions)}?" if suggestions else ""
    raise KeyError(f"Gene {gene_query} was not found.{hint}")


def _normalized_obs_column(column: str) -> np.ndarray:
    with OBS_COLUMN_CACHE_LOCK:
        cached = OBS_COLUMN_CACHE.get(column)
        if cached is not None:
            return cached

    adata = load_gene_expression_adata()
    values = normalize_series(adata.obs[column]).to_numpy()
    with OBS_COLUMN_CACHE_LOCK:
        OBS_COLUMN_CACHE[column] = values
    return values


def _lowercase_obs_column(column: str) -> np.ndarray:
    key = f"{column}:lower"
    with OBS_COLUMN_CACHE_LOCK:
        cached = OBS_COLUMN_CACHE.get(key)
    if cached is None:
        cached = pd.Series(_normalized_obs_column(column)).str.lower().to_numpy()
        with OBS_COLUMN_CACHE_LOCK:
            OBS_COLUMN_CACHE[key] = cached
    return cached


def _gene_vector_cache_path(cache_key: str) -> str:
    stat = os.stat(GENE_EXPRESSION_DATA_PATH)
    # Tie the cache to the source file so a regenerated h5ad invalidates it.
    version = f"{stat.st_mtime_ns}_{stat.st_size}"
    directory = os.path.join(GENE_VECTOR_CACHE_DIR, version)
    os.makedirs(directory, exist_ok=True)
    return os.path.join(directory, f"{cache_key}.npy")


def _raw_gene_expression_vector(adata: ExpressionStore, gene_idx: int) -> np.ndarray:
    # Scans the whole raw CSR matrix (~30 s); results are cached on disk.
    counts = adata.column("raw/X", gene_idx).astype(np.float64)

    if "total_counts" in adata.obs.columns:
        totals = adata.obs["total_counts"].to_numpy(dtype=np.float64)
        positive_totals = totals[totals > 0]
        target_sum = float(np.median(positive_totals)) if positive_totals.size else 1.0
        with np.errstate(divide="ignore", invalid="ignore"):
            counts = np.where(totals > 0, counts / totals * target_sum, 0.0)

    return np.log1p(counts)


def _gene_expression_vector(adata: ExpressionStore, gene_idx: int, source: str = "main") -> np.ndarray:
    cache_key = f"{source}_{gene_idx}"
    with GENE_VECTOR_CACHE_LOCK:
        cached = GENE_VECTOR_CACHE.get(cache_key)
        if cached is not None:
            return cached

    cache_path = _gene_vector_cache_path(cache_key)
    values = None
    if os.path.isfile(cache_path):
        try:
            values = np.load(cache_path, mmap_mode="r")
        except (ValueError, OSError):
            values = None

    if values is None:
        if source == "raw":
            values = _raw_gene_expression_vector(adata, gene_idx)
        else:
            values = adata.column(adata.main_matrix, gene_idx)
        values = np.nan_to_num(values.astype(np.float32), nan=0.0, posinf=0.0, neginf=0.0)
        temp_path = f"{cache_path}.{os.getpid()}.tmp.npy"
        try:
            np.save(temp_path, values)
            os.replace(temp_path, cache_path)
        except OSError:
            if os.path.exists(temp_path):
                os.remove(temp_path)

    with GENE_VECTOR_CACHE_LOCK:
        GENE_VECTOR_CACHE[cache_key] = values
        while len(GENE_VECTOR_CACHE) > GENE_VECTOR_CACHE_SIZE:
            GENE_VECTOR_CACHE.pop(next(iter(GENE_VECTOR_CACHE)))
    return values


def _selection_mask(n_obs: int, cell_type: str | None = None, disease: str | None = None) -> np.ndarray:
    adata = load_gene_expression_adata()
    mask = np.ones(n_obs, dtype=bool)
    selected_cell_types = _parse_multi_select_values(cell_type)
    selected_diseases = _parse_multi_select_values(disease)

    if selected_cell_types and "cell_type2" in adata.obs.columns:
        allowed = {value.strip().lower() for value in selected_cell_types}
        mask &= np.isin(_lowercase_obs_column("cell_type2"), list(allowed))

    if selected_diseases and "diagnosis" in adata.obs.columns:
        allowed = {value.strip().lower() for value in selected_diseases}
        mask &= np.isin(_lowercase_obs_column("diagnosis"), list(allowed))

    return mask


def _dataframe_for_gene_expression(adata: ExpressionStore, gene_idx: int, cell_type: str | None = None, disease: str | None = None, source: str = "main") -> pd.DataFrame:
    mask = _selection_mask(adata.n_obs, cell_type=cell_type, disease=disease)
    values = _gene_expression_vector(adata, gene_idx, source=source)[mask]

    frame = pd.DataFrame({"expression": values})
    frame["cell_type"] = _normalized_obs_column("cell_type2")[mask] if "cell_type2" in adata.obs.columns else "All cells"
    frame["diagnosis"] = _normalized_obs_column("diagnosis")[mask] if "diagnosis" in adata.obs.columns else "All"
    return frame[frame["expression"] >= 0].reset_index(drop=True)


def _encode_plot_png(fig) -> str:
    buffer = io.BytesIO()
    fig.tight_layout()
    fig.savefig(buffer, format="png", dpi=110, bbox_inches="tight")
    plt.close(fig)
    return "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode("utf-8")


def _expression_group_orders(frame: pd.DataFrame) -> tuple[list[str], list[str]]:
    cell_order = sorted(frame["cell_type"].dropna().astype(str).unique().tolist())
    disease_order = sorted(frame["diagnosis"].dropna().astype(str).unique().tolist())
    return cell_order, disease_order


# Rendering every cell is the dominant cost; a per-group cap keeps distributions faithful.
def _downsample_expression_frame(frame: pd.DataFrame, max_per_group: int = GENE_PLOT_MAX_PER_GROUP) -> pd.DataFrame:
    if frame.empty or len(frame) <= max_per_group:
        return frame
    groups = max(1, frame.groupby(["cell_type", "diagnosis"], observed=True, sort=False).ngroups)
    per_group = max(200, min(max_per_group, GENE_PLOT_MAX_POINTS // groups))
    shuffled = frame.iloc[np.random.default_rng(0).permutation(len(frame))]
    keep = shuffled.groupby(["cell_type", "diagnosis"], observed=True, sort=False).cumcount() < per_group
    return shuffled[keep].reset_index(drop=True)


def _get_expression_chart_image(gene_symbol: str, frame: pd.DataFrame) -> str:
    sns.set_theme(style="whitegrid", font_scale=1.4)

    if frame.empty:
        fig, ax = plt.subplots(figsize=(8, 4))
        ax.text(0.5, 0.5, "No cells with the current filters", ha="center", va="center")
        ax.axis("off")
        return _encode_plot_png(fig)

    cell_order, disease_order = _expression_group_orders(frame)
    col_wrap = min(4, max(1, len(cell_order)))
    grid = sns.displot(
        data=frame,
        x="expression",
        col="cell_type",
        col_order=cell_order,
        col_wrap=col_wrap,
        hue="diagnosis",
        hue_order=disease_order,
        palette="Set2",
        bins=25,
        element="step",
        common_norm=False,
        stat="count",
        height=3.6,
        aspect=1.4,
        facet_kws={"sharex": False, "sharey": False},
    )
    grid.set_axis_labels("Expression", "Cells", fontsize=18)
    grid.set_titles(col_template="{col_name}", size=16)
    if grid.legend is not None:
        sns.move_legend(grid, "center left", bbox_to_anchor=(1.0, 0.5), frameon=True)
    grid.figure.subplots_adjust(right=0.86, hspace=0.45, wspace=0.25)
    grid.figure.suptitle(f"{gene_symbol} expression by cell type and diagnosis", y=1.03, fontsize=24)
    return _encode_plot_png(grid.figure)


def _parse_multi_select_values(raw_value: str | None) -> list[str]:
    if raw_value is None:
        return []
    return [item.strip() for item in str(raw_value).split(",") if item.strip()]


def _get_violin_chart_image(gene_symbol: str, frame: pd.DataFrame) -> str:
    sns.set_theme(style="whitegrid", font_scale=1.8)
    if frame.empty:
        fig, ax = plt.subplots(figsize=(16, 6))
        ax.text(0.5, 0.5, "No cells with the current filters", ha="center", va="center")
        ax.axis("off")
        return _encode_plot_png(fig)

    cell_order, disease_order = _expression_group_orders(frame)
    grid = sns.catplot(
        data=frame,
        kind="violin",
        x="cell_type",
        y="expression",
        order=cell_order,
        row="diagnosis",
        row_order=disease_order,
        hue="cell_type",
        hue_order=cell_order,
        legend=False,
        inner="box",
        cut=0,
        density_norm="width",
        palette="Set2",
        height=4.4,
        aspect=max(2.6, len(cell_order) * 0.38),
        sharex=True,
        sharey=True,
    )
    grid.set_axis_labels("Cell type", "Expression", fontsize=22)
    grid.set_titles(row_template="{row_name}", size=20)
    for ax in grid.axes.flat:
        ax.tick_params(axis="x", rotation=45)
        for label in ax.get_xticklabels():
            label.set_ha("right")
    grid.figure.suptitle(f"{gene_symbol} expression by cell type, split by diagnosis", y=1.01, fontsize=30)
    return _encode_plot_png(grid.figure)


def get_gene_expression(gene_query: str, cell_type: str | None = None, disease: str | None = None, plot_type: str = "histogram"):
    adata = load_gene_expression_adata()
    gene_idx, gene_symbol, source = _resolve_gene_index(adata, gene_query)
    frame = _dataframe_for_gene_expression(adata, gene_idx, cell_type=cell_type, disease=disease, source=source)
    values = frame["expression"].to_numpy()

    cache_key = (gene_symbol, plot_type, cell_type or "", disease or "")
    with GENE_PLOT_CACHE_LOCK:
        image_base64 = GENE_PLOT_CACHE.get(cache_key)
    if image_base64 is None:
        plot_frame = _downsample_expression_frame(frame)
        if plot_type == "violin":
            image_base64 = _get_violin_chart_image(gene_symbol, plot_frame)
        else:
            image_base64 = _get_expression_chart_image(gene_symbol, plot_frame)
        with GENE_PLOT_CACHE_LOCK:
            GENE_PLOT_CACHE[cache_key] = image_base64
            while len(GENE_PLOT_CACHE) > GENE_PLOT_CACHE_SIZE:
                GENE_PLOT_CACHE.pop(next(iter(GENE_PLOT_CACHE)))

    return {
        "gene": gene_symbol,
        "cell_type_column": "cell_type2" if "cell_type2" in adata.obs.columns else None,
        "cell_type": cell_type or "All",
        "selected_cells": int(values.size),
        "nonzero_cells": int(np.count_nonzero(values > 0)),
        "mean_expression": float(np.mean(values)) if values.size else 0.0,
        "median_expression": float(np.median(values)) if values.size else 0.0,
        "max_expression": float(np.max(values)) if values.size else 0.0,
        "expression_values": values.tolist()[:5000],
        "image_base64": image_base64,
    }


def get_gene_violin(gene_query: str, cell_types: str | None = None, disease: str | None = None):
    adata = load_gene_expression_adata()
    gene_query = (gene_query or "").strip()
    if not gene_query:
        raise ValueError("A gene name is required")

    gene_idx, gene_symbol, source = _resolve_gene_index(adata, gene_query)
    effective_column = "cell_type2" if "cell_type2" in adata.obs.columns else None
    selected_categories = _parse_multi_select_values(cell_types)
    selected_diseases = _parse_multi_select_values(disease)

    frame = _dataframe_for_gene_expression(adata, gene_idx, cell_type=cell_types, disease=disease, source=source)
    groups = [
        {"name": group_name, "values": group["expression"].tolist()}
        for group_name, group in frame.groupby(["cell_type", "diagnosis"], dropna=False)
    ]

    return {
        "gene": gene_symbol,
        "cell_type_column": effective_column,
        "cell_type": ", ".join(selected_categories) if selected_categories else "All",
        "disease": ", ".join(selected_diseases) if selected_diseases else "All",
        "series": groups,
        "image_base64": _get_violin_chart_image(gene_symbol, frame),
    }


def parse_query(query: str):
    text = (query or "").strip().lower()
    if "count" in text and ("dataset" in text or "type" in text):
        if "dataset" in text:
            return {"action": "count", "column": "dataset"}
        if "type" in text:
            return {"action": "count", "column": "Type"}
    if "plot" in text or "visualize" in text or "umap" in text:
        color_col = "dataset" if "dataset" in text else "Type" if "type" in text else "cell_type2"
        return {"action": "plot", "column": color_col}
    if "show" in text or "list" in text or "search" in text:
        if "dataset" in text:
            return {"action": "search", "column": "dataset"}
        if "type" in text:
            return {"action": "search", "column": "Type"}
        if "cell" in text:
            return {"action": "search", "column": "Cell"}
    if "summary" in text or "overview" in text:
        return {"action": "summary"}
    return {"action": "summary"}


ASSISTANT_MODEL = os.environ.get("ASSISTANT_MODEL", "gpt-5.4-nano")
ASSISTANT_OUT_OF_SCOPE = "This question is beyond the knowledge scope of LiverHomo."
ASSISTANT_MAX_INPUT_TOKENS = 1000
ASSISTANT_MAX_OUTPUT_TOKENS = 2500
ASSISTANT_MAX_ROUNDS = 5
ASSISTANT_HISTORY_MESSAGES = 6
ASSISTANT_HISTORY_MESSAGE_CHARS = 1500
ASSISTANT_TOOL_RESULT_CHARS = 12_000
COMPOSITION_PATH = os.path.join(app.root_path, "static", "composition.json")
DEG_INDEX = DegIndex(DEG_ROOT)
_NAME_RESOLVER: NameResolver | None = None

_DISEASE_ARG = {"type": "string", "description": "Disease in the user's own words, e.g. 'HCC', 'NASH', 'primary biliary cholangitis'. Comma-separate several."}
_CELL_TYPE_ARG = {"type": "string", "description": "Cell type in the user's own words, e.g. 'hepatocytes', 'Kupffer cells', 'CD8 T cells'. Groups such as 'T cells' expand to all subsets. Comma-separate several."}
_GENE_ARG = {"type": "string", "description": "Official gene symbol, e.g. 'SPP1'."}


def _assistant_tool(name: str, description: str, properties: dict[str, Any], required: list[str]) -> dict[str, Any]:
    return {
        "type": "function",
        "function": {
            "name": name,
            "description": description,
            "parameters": {"type": "object", "properties": properties, "required": required, "additionalProperties": False},
        },
    }


ASSISTANT_TOOLS = [
    _assistant_tool(
        "search_gene_deg",
        "Differential expression results for ONE gene: every disease and cell type where it is significantly up "
        "or down, with effect size (de_coef) and p-value, plus combinations that were tested but not significant. "
        "Use for 'Is gene X changed/up/down in ...?' and 'Where is gene X differentially expressed?'.",
        {"gene": _GENE_ARG, "disease": _DISEASE_ARG, "cell_type": _CELL_TYPE_ARG},
        ["gene"],
    ),
    _assistant_tool(
        "top_deg_genes",
        "Ranked list of significant differentially expressed genes for a disease, optionally restricted to cell "
        "types and direction. Use for 'top genes', 'list genes up/down in ...'.",
        {
            "disease": _DISEASE_ARG,
            "cell_type": _CELL_TYPE_ARG,
            "direction": {"type": "string", "enum": ["up", "down", "both"]},
            "sort_by": {"type": "string", "enum": ["pvalue", "effect"], "description": "pvalue (default) or absolute de_coef."},
            "limit": {"type": "integer", "minimum": 1, "maximum": 30},
        },
        ["disease"],
    ),
    _assistant_tool(
        "gene_expression_summary",
        "Expression level of one gene per disease and cell type: percent of cells expressing it and mean "
        "log-normalized expression. Use for 'Is gene X expressed in ...?' or 'Which cells express gene X?'. "
        "The first request for a gene can take up to a minute.",
        {"gene": _GENE_ARG, "disease": _DISEASE_ARG, "cell_type": _CELL_TYPE_ARG},
        ["gene"],
    ),
    _assistant_tool(
        "cell_counts",
        "Number of cells per disease and cell type in the atlas.",
        {"disease": _DISEASE_ARG, "cell_type": _CELL_TYPE_ARG},
        [],
    ),
    _assistant_tool(
        "atlas_overview",
        "Diseases (with full names), cell types, total cells and genes, and which analyses are available.",
        {},
        [],
    ),
    _assistant_tool(
        "generate_visualization",
        "Show a figure to the user. gene_violin: expression of a gene by cell type (optionally filtered). "
        "gene_expression_histogram: distribution of a gene's expression. volcano: DEG volcano plot for a disease. "
        "umap_metadata: whole-atlas UMAP coloured by a metadata column (cannot be filtered). "
        "cell_interaction_heatmap: cell-cell interaction counts for a disease. "
        "cell_interaction_pair_histogram: interactions between two cell types across diseases. "
        "ppi_network: protein-protein interaction neighbourhood of a gene in one disease and cell type.",
        {
            "chart_type": {"type": "string", "enum": [
                "gene_violin", "gene_expression_histogram", "volcano", "umap_metadata",
                "cell_interaction_heatmap", "cell_interaction_pair_histogram", "ppi_network",
            ]},
            "gene": _GENE_ARG,
            "disease": _DISEASE_ARG,
            "cell_type": _CELL_TYPE_ARG,
            "column": {"type": "string", "enum": list(UMAP_METADATA_COLUMNS)},
            "cell_type_a": {"type": "string"},
            "cell_type_b": {"type": "string"},
        },
        ["chart_type"],
    ),
]

ASSISTANT_SYSTEM_PROMPT = """You are the assistant of LiverHomo, a single-cell RNA-seq atlas of the human liver \
(about 1.16 million cells, 26 cell types, 11 disease groups plus Healthy). Users are biologists and clinicians who \
mostly ask whether their genes of interest are changed or expressed in particular diseases and cell types.

Rules:
1. Every statement about atlas data (DEG results, expression, cell counts, figures) must come from a tool result in \
this conversation. Never invent genes, numbers or results. Call the tools you need before answering.
2. Pass diseases and cell types in the user's own words; the tools understand synonyms (NASH -> MASH, Kupffer cells \
-> Macrophage, T cells -> all T-cell subsets). If a tool returns an error listing valid names, retry with a valid \
name or ask the user to choose.
3. DEG results contain only significant genes. If a gene was tested in a disease/cell type but is absent, say it is \
not significantly changed there. If that combination was not analysed, say so. Every disease group is compared with \
Healthy (Healthy itself has no DEG results). "up"/"down" and a positive/negative de_coef mean higher/lower in the \
disease than in Healthy; de_coef is a log-scale effect size and de_pval the p-value.
4. You may add one or two sentences of general biological background (for example what a gene encodes) when it \
helps. Such sentences MUST start with the exact label "Background (general knowledge):", for example \
"Background (general knowledge): SPP1 encodes osteopontin." Never present general knowledge as atlas data. When asked what \
a gene does or what its function is, call search_gene_deg and answer in exactly this shape:
   Background (general knowledge): <one or two sentences on what the gene encodes and does>
   In LiverHomo: <where it is significantly up/down, with the strongest examples and totals>
5. Answer directly first (yes/no or the list), then the key numbers: cell type, direction, de_coef, p-value. Use \
the atlas names for diseases and cell types. For long results show the most significant ones and give the total.
6. When the user asks to see, show, plot or visualize something, call generate_visualization and describe the \
figure briefly.
7. No tables or headings; **bold** only for the direct answer or gene names. Keep answers under 200 words unless the user asks for more.
8. If the question is not about LiverHomo data or the biology of the genes, cell types and diseases in it (for \
example treatment, diagnosis or other medical advice, news, weather, programming), reply with exactly this \
sentence and nothing else: """ + ASSISTANT_OUT_OF_SCOPE


def get_name_resolver() -> NameResolver:
    global _NAME_RESOLVER
    if _NAME_RESOLVER is None:
        composition = load_composition(COMPOSITION_PATH)
        diseases = {short_disease_name(name): name for name in composition}
        cell_types = sorted({cell_type for counts in composition.values() for cell_type in counts})
        _NAME_RESOLVER = NameResolver(diseases, cell_types)
    return _NAME_RESOLVER


def _resolve_diseases(text: str | None) -> list[str]:
    resolver = get_name_resolver()
    return [resolver.disease(part) for part in re.split(r",|;", text or "") if part.strip()]


def _gene_expression_summary(gene: str, diseases: list[str], cell_types: list[str]) -> dict[str, Any]:
    adata = load_gene_expression_adata()
    gene_idx, symbol, source = _resolve_gene_index(adata, gene)
    frame = _dataframe_for_gene_expression(
        adata, gene_idx, cell_type=",".join(cell_types) or None, disease=",".join(diseases) or None, source=source,
    )
    if frame.empty:
        return {"gene": symbol, "groups": [], "note": "No cells match that selection."}
    frame["expressing"] = frame["expression"] > 0
    grouped = frame.groupby(["diagnosis", "cell_type"], observed=True).agg(
        cells=("expression", "size"), pct_expressing=("expressing", "mean"), mean_expression=("expression", "mean"),
    ).reset_index()
    grouped = grouped[grouped["cells"] >= 20].sort_values("mean_expression", ascending=False)
    groups = [
        {
            "disease": row.diagnosis, "cell_type": row.cell_type, "cells": int(row.cells),
            "pct_expressing": round(100 * float(row.pct_expressing), 1), "mean_expression": round(float(row.mean_expression), 3),
        }
        for row in grouped.head(30).itertuples(index=False)
    ]
    return {
        "gene": symbol,
        "values": "log-normalized expression; groups with fewer than 20 cells omitted; sorted by mean expression",
        "groups_total": int(len(grouped)),
        "groups": groups,
    }


def _cell_counts(diseases: list[str], cell_types: list[str]) -> dict[str, Any]:
    composition = load_composition(COMPOSITION_PATH)
    rows = []
    for full_name, counts in composition.items():
        disease = short_disease_name(full_name)
        if diseases and disease not in diseases:
            continue
        for cell_type, count in counts.items():
            if cell_types and cell_type not in cell_types:
                continue
            rows.append({"disease": disease, "cell_type": cell_type, "cells": int(count)})
    rows.sort(key=lambda row: -row["cells"])
    return {"total_cells": sum(row["cells"] for row in rows), "rows": rows[:60], "rows_total": len(rows)}


def _atlas_overview() -> dict[str, Any]:
    resolver = get_name_resolver()
    composition = load_composition(COMPOSITION_PATH)
    deg_diseases = sorted({disease for disease, _ in DEG_INDEX.tested})
    adata = load_gene_expression_adata()
    return {
        "total_cells": sum(sum(counts.values()) for counts in composition.values()),
        "diseases": [
            {"name": short, "full_name": full, "cells": sum(composition[full].values()), "has_deg_results": short in deg_diseases}
            for short, full in resolver.diseases.items()
        ],
        "cell_types": resolver.cell_types,
        "genes_with_expression_data": int(adata.raw.var_names.size if adata.raw is not None else adata.n_vars),
        "available": [
            "differential expression (DEG) per disease and cell type",
            "gene expression by disease and cell type",
            "cell counts", "volcano plots", "UMAP images", "cell-cell interaction counts",
            "protein-protein interaction networks per disease and cell type",
        ],
    }


def get_assistant_data(name: str, arguments: dict[str, Any]) -> dict[str, Any]:
    """Run a data tool. Raises LookupError/KeyError/ValueError with a message the model can act on."""
    resolver = get_name_resolver()
    diseases = _resolve_diseases(arguments.get("disease"))
    cell_types = resolver.cell_types_for(arguments.get("cell_type"))
    gene = str(arguments.get("gene") or "").strip()

    if name == "search_gene_deg":
        if not gene:
            raise ValueError("gene is required")
        _resolve_gene_index(load_gene_expression_adata(), gene)  # rejects unknown genes with suggestions
        return DEG_INDEX.gene(gene, diseases, cell_types)
    if name == "top_deg_genes":
        if len(diseases) != 1:
            raise ValueError("Exactly one disease is required for top_deg_genes.")
        direction = arguments.get("direction") if arguments.get("direction") in ("up", "down") else None
        limit = min(max(int(arguments.get("limit") or 10), 1), 30)
        sort = "effect" if arguments.get("sort_by") == "effect" else "pvalue"
        return DEG_INDEX.top(diseases[0], cell_types, direction, limit, sort)
    if name == "gene_expression_summary":
        if not gene:
            raise ValueError("gene is required")
        return _gene_expression_summary(gene, diseases, cell_types)
    if name == "cell_counts":
        return _cell_counts(diseases, cell_types)
    if name == "atlas_overview":
        return _atlas_overview()
    raise ValueError(f"Unknown tool {name}")


def _encode_image_file_png(path: str) -> str:
    with open(path, "rb") as handle:
        return base64.b64encode(handle.read()).decode("ascii")


def _render_network_png(nodes: list[dict[str, Any]], edges: list[dict[str, Any]], title: str) -> str:
    sns.set_theme(style="white")
    fig, ax = plt.subplots(figsize=(8, 7))
    positions = {node["id"]: (node["x"], node["y"]) for node in nodes}
    for edge in edges:
        source_pos = positions.get(edge.get("source"))
        target_pos = positions.get(edge.get("target"))
        if source_pos and target_pos:
            ax.plot([source_pos[0], target_pos[0]], [source_pos[1], target_pos[1]], color="#b7c4cf", linewidth=0.8, zorder=1)
    ax.scatter(
        [node["x"] for node in nodes],
        [node["y"] for node in nodes],
        s=[30 + 12 * node.get("degree", 0) for node in nodes],
        color="#c1272d",
        edgecolor="white",
        linewidth=0.5,
        zorder=2,
    )
    for node in nodes:
        ax.annotate(node["id"], (node["x"], node["y"]), fontsize=8, xytext=(3, 3), textcoords="offset points")
    ax.set_title(title)
    ax.axis("off")
    return _encode_plot_png(fig)


ASSISTANT_VISUALIZATION_LINKS = {
    "gene_expression_histogram": ("/atlas#expression-view", "Open Gene Expression"),
    "gene_violin": ("/atlas#expression-view", "Open Gene Expression"),
    "umap_metadata": ("/atlas#explore-view", "Open Atlas Explorer"),
    "volcano": ("/atlas#deg-view", "Open DEGs and Volcano Plots"),
    "cell_interaction_heatmap": ("/graph#interaction-view", "Open Cell-Cell Interactions"),
    "cell_interaction_pair_histogram": ("/graph#interaction-view", "Open Cell-Cell Interactions"),
    "ppi_network": ("/graph#graph-view", "Open Protein Interaction Network"),
}


def _assistant_visualization_result(chart_type: str, title: str, image_base64: str) -> dict[str, str]:
    destination_url, destination_label = ASSISTANT_VISUALIZATION_LINKS[chart_type]
    return {
        "chart_type": chart_type,
        "title": title,
        "image_base64": image_base64,
        "destination_url": destination_url,
        "destination_label": destination_label,
    }


def get_assistant_visualization(arguments: dict[str, Any]) -> dict[str, Any]:
    """Build a chart. Raises LookupError/KeyError/ValueError with a message the model can act on."""
    resolver = get_name_resolver()
    chart_type = str(arguments.get("chart_type") or "").strip()
    gene = str(arguments.get("gene") or "").strip()
    diseases = _resolve_diseases(arguments.get("disease"))
    cell_types = resolver.cell_types_for(arguments.get("cell_type"))
    disease_arg = ",".join(diseases) or None
    cell_type_arg = ",".join(cell_types) or None

    if chart_type in ("gene_expression_histogram", "gene_violin", "ppi_network") and not gene:
        raise ValueError(f"gene is required for {chart_type}")
    if chart_type == "gene_expression_histogram":
        payload = get_gene_expression(gene, cell_type=cell_type_arg, disease=disease_arg, plot_type="histogram")
        return _assistant_visualization_result(chart_type, f"{payload['gene']} expression distribution", payload["image_base64"])
    if chart_type == "gene_violin":
        payload = get_gene_violin(gene, cell_types=cell_type_arg, disease=disease_arg)
        return _assistant_visualization_result(chart_type, f"{payload['gene']} expression by cell type", payload["image_base64"])
    if chart_type == "umap_metadata":
        column = str(arguments.get("column") or "").strip()
        image_path = get_static_umap_image_path(column)
        if image_path is None or not os.path.isfile(image_path):
            raise ValueError(f"No UMAP image for column '{column}'. Available: {', '.join(get_static_umap_metadata_columns())}")
        return _assistant_visualization_result(chart_type, f"UMAP colored by {column}", _encode_image_file_png(image_path))
    if chart_type in ("volcano", "cell_interaction_heatmap"):
        if len(diseases) != 1:
            raise ValueError(f"Exactly one disease is required for {chart_type}.")
        if chart_type == "volcano":
            path, title = get_volcano_path(diseases[0]), f"Volcano plot: {diseases[0]}"
        else:
            path, title = get_cell_interaction_path(diseases[0], kind="heatmap"), f"Cell-cell interaction heatmap: {diseases[0]}"
        if path is None:
            raise LookupError(f"No {chart_type} figure exists for {diseases[0]}.")
        return _assistant_visualization_result(chart_type, title, _encode_image_file_png(path))
    if chart_type == "cell_interaction_pair_histogram":
        pair = [resolver.cell_types_for(arguments.get(key)) for key in ("cell_type_a", "cell_type_b")]
        if any(len(names) != 1 for names in pair):
            raise ValueError("cell_type_a and cell_type_b must each name exactly one cell type.")
        payload = get_cell_interaction_pair_counts(pair[0][0], pair[1][0])
        if not payload["rows"]:
            raise LookupError("No interaction counts found for that pair.")
        return _assistant_visualization_result(chart_type, f"{pair[0][0]} ↔ {pair[1][0]} interactions by disease", payload["image_base64"])
    if chart_type == "ppi_network":
        if len(diseases) != 1 or len(cell_types) != 1:
            raise ValueError("ppi_network needs exactly one disease and one cell type.")
        graph_payload = get_graph_payload(disease=diseases[0], cell_type=cell_types[0], limit=1, gene_list=gene, max_edges=GRAPH_LAYOUT_EDGE_LIMIT)
        graphs = graph_payload.get("graphs") or []
        if not graphs or not graphs[0].get("nodes"):
            raise LookupError(f"{gene} is not in the {diseases[0]} / {cell_types[0]} interaction network.")
        network = build_graph_network({"nodes": graphs[0]["nodes"], "edges": graphs[0]["edges"]}, layout="kamada_kawai")
        title = f"{gene} PPI network — {diseases[0]} / {cell_types[0]}"
        return _assistant_visualization_result(chart_type, title, _render_network_png(network["nodes"], network["edges"], title))
    raise ValueError(f"Unknown chart_type '{chart_type}'")


def limit_assistant_input(query: str) -> str:
    encoding = tiktoken.get_encoding("o200k_base")
    tokens = encoding.encode(query)
    return encoding.decode(tokens[:ASSISTANT_MAX_INPUT_TOKENS])


def _assistant_history(history: Any) -> list[dict[str, str]]:
    if not isinstance(history, list):
        return []
    messages = []
    for item in history[-ASSISTANT_HISTORY_MESSAGES:]:
        if isinstance(item, dict) and item.get("role") in ("user", "assistant") and isinstance(item.get("content"), str):
            messages.append({"role": item["role"], "content": item["content"][:ASSISTANT_HISTORY_MESSAGE_CHARS]})
    return messages


def _run_assistant_tool(name: str, raw_arguments: str) -> tuple[str, dict[str, Any] | None]:
    """Returns (tool message content, visualization or None)."""
    try:
        arguments = json.loads(raw_arguments or "{}")
        if not isinstance(arguments, dict):
            raise ValueError("Tool arguments must be a JSON object.")
        if name == "generate_visualization":
            result = get_assistant_visualization(arguments)
            summary = {key: value for key, value in result.items() if key != "image_base64"}
            return json.dumps({**summary, "status": "Chart generated and shown to the user."}), result
        content = json.dumps(get_assistant_data(name, arguments), default=str)
        if len(content) > ASSISTANT_TOOL_RESULT_CHARS:
            content = content[:ASSISTANT_TOOL_RESULT_CHARS] + " ...[truncated]"
        return content, None
    except (LookupError, ValueError, FileNotFoundError, json.JSONDecodeError) as exc:
        message = exc.args[0] if exc.args else str(exc)
        return json.dumps({"error": str(message)}), None
    except Exception:  # keep the conversation going; details go to the log
        app.logger.exception("Assistant tool %s failed", name)
        return json.dumps({"error": "Internal error while running this tool."}), None


def call_openai(query: str, history: Any = None):
    query = limit_assistant_input(query.strip())
    if not OPENAI_API_KEY or not query:
        return {"error": {"message": "OpenAI API key is not configured or query is empty.", "type": "configuration_error"}}, 503

    messages: list[dict[str, Any]] = [
        {"role": "system", "content": ASSISTANT_SYSTEM_PROMPT},
        *_assistant_history(history),
        {"role": "user", "content": query},
    ]
    visualization: dict[str, Any] | None = None
    tools_used: list[str] = []
    usage_total = 0
    try:
        for round_index in range(ASSISTANT_MAX_ROUNDS):
            final_round = round_index == ASSISTANT_MAX_ROUNDS - 1
            payload = {
                "model": ASSISTANT_MODEL,
                "messages": messages,
                "tools": ASSISTANT_TOOLS,
                "tool_choice": "none" if final_round else "auto",
                "reasoning_effort": "none",
                "max_completion_tokens": ASSISTANT_MAX_OUTPUT_TOKENS,
            }
            req = urllib_request.Request(
                "https://api.openai.com/v1/chat/completions",
                headers={"Authorization": f"Bearer {OPENAI_API_KEY}", "Content-Type": "application/json"},
                data=json.dumps(payload).encode("utf-8"),
            )
            with urllib_request.urlopen(req, timeout=60) as resp:
                package = json.loads(resp.read().decode("utf-8"))
            usage_total += int(package.get("usage", {}).get("total_tokens", 0))
            message = package["choices"][0]["message"]
            tool_calls = message.get("tool_calls") or []
            if not tool_calls:
                answer = str(message.get("content") or "").strip()
                app.logger.info("Assistant query=%r tools=%s tokens=%d", query[:200], tools_used, usage_total)
                if answer:
                    return {"answer": answer, "visualization": visualization}, 200
                return {"error": {"message": "The assistant did not return an answer. Please rephrase the question.", "type": "response_error"}}, 502

            messages.append({key: message[key] for key in ("role", "content", "tool_calls") if key in message})
            for tool_call in tool_calls:
                name = tool_call.get("function", {}).get("name", "")
                tools_used.append(name)
                content, chart = _run_assistant_tool(name, tool_call.get("function", {}).get("arguments", ""))
                if chart is not None:
                    visualization = chart
                messages.append({"role": "tool", "tool_call_id": tool_call["id"], "content": content})
    except urllib_error.HTTPError as exc:
        try:
            error_package = json.loads(exc.read().decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            error_package = {"error": {"message": str(exc.reason), "type": "http_error"}}
        app.logger.error("OpenAI API request failed: status=%s response=%s", exc.code, error_package)
        return {"error": {"message": "The AI service returned an error. Please try again later.", "type": "upstream_error"}}, 502
    except (urllib_error.URLError, TimeoutError) as exc:
        app.logger.error("OpenAI API unreachable: %s", exc)
        return {"error": {"message": "The AI service could not be reached. Please try again later.", "type": "network_error"}}, 502
    except (json.JSONDecodeError, KeyError, IndexError) as exc:
        app.logger.error("Unexpected OpenAI response: %r", exc)
        return {"error": {"message": "Unexpected response from the AI service.", "type": "response_error"}}, 502
    return {"error": {"message": "The assistant did not return an answer. Please rephrase the question.", "type": "response_error"}}, 502


@app.route("/")
def index():
    unload_graph_data()
    return render_template("index.html")


@app.route("/atlas")
def atlas_page():
    unload_graph_data()
    return render_template("atlas.html", active_tab="atlas")


@app.route("/graph")
def graph_page():
    return render_template("graph.html", active_tab="graph")


@app.route("/resources")
def resources_page():
    unload_graph_data()
    return render_template("resources.html", active_tab="resources")


@app.route("/metadata_download")
def metadata_download():
    dataset_path = os.path.join(app.root_path, "adata_reprocessed_gzip.h5ad")
    return send_file(
        dataset_path,
        as_attachment=True,
        download_name="adata_reprocessed_gzip.h5ad",
    )


@app.route("/contact")
def contact_page():
    return render_template("contact.html", active_tab="contact")


@app.route("/api/contact", methods=["POST"])
def api_contact():
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("name") or "").strip()
    sender_email = str(payload.get("email") or "").strip()
    message = str(payload.get("message") or "").strip()

    if not name or not sender_email or not message:
        return jsonify({"error": "Name, email, and message are required."}), 400
    if len(name) > 120 or len(sender_email) > 254 or len(message) > 5_000:
        return jsonify({"error": "Message is too long."}), 400
    if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", sender_email):
        return jsonify({"error": "Enter a valid email address."}), 400
    if not all([SMTP_HOST, SMTP_USERNAME, SMTP_PASSWORD, SMTP_FROM_EMAIL]):
        return jsonify({"error": "Contact email is not configured."}), 503

    email = EmailMessage()
    email["Subject"] = f"LiverHomo/DCP contact message from {name}"
    email["From"] = SMTP_FROM_EMAIL
    email["To"] = CONTACT_RECIPIENT
    email["Reply-To"] = sender_email
    email.set_content(f"Name: {name}\nEmail: {sender_email}\n\nMessage:\n{message}")

    try:
        with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=20) as smtp:
            smtp.starttls()
            smtp.login(SMTP_USERNAME, SMTP_PASSWORD)
            smtp.send_message(email)
    except (OSError, smtplib.SMTPException):
        return jsonify({"error": "Message could not be sent. Please try again later."}), 502
    return jsonify({"message": "Your message has been sent."})


@app.route("/api/summary")
def api_summary():
    with BACKGROUND_DATA_LOCK:
        cached = BACKGROUND_DATA["summary"]
        if cached is not None:
            return jsonify(cached)
        loading = BACKGROUND_DATA["loading"]

    if not loading:
        start_background_data_load()

    filter_catalog = load_filter_catalog()
    return jsonify({
        "loading": True,
        "cells": 0,
        "genes": 0,
        "obs_columns": 0,
        "sample_columns": [],
        "dataset_counts": {},
        "type_counts": {},
        "cell_type_column": None,
        "cell_type_options": filter_catalog["metadata"]["cell_types"],
        "disease_column": None,
        "disease_options": filter_catalog["metadata"]["diseases"],
        "umap_columns": get_static_umap_metadata_columns(),
    })


@app.route("/api/search")
def api_search():
    column = request.args.get("column", "").strip()
    value = request.args.get("value", "").strip()
    limit = int(request.args.get("limit", 50))
    cell_type = request.args.get("cell_type", "").strip()
    disease = request.args.get("disease", "").strip()
    sort_by = request.args.get("sort_by", "").strip()
    sort_order = request.args.get("sort_order", "asc").strip()
    records = search_rows(
        column or None,
        value or None,
        limit=limit,
        cell_type=cell_type or None,
        disease=disease or None,
        sort_by=sort_by or None,
        sort_order=sort_order,
    )
    return jsonify({"records": records, "plot": None, "count": len(records)})


@app.route("/api/umap_metadata")
def api_umap_metadata():
    column = request.args.get("column", "").strip()
    cell_type = request.args.get("cell_type", "").strip()
    disease = request.args.get("disease", "").strip()
    if not column:
        return jsonify({"error": "No column provided for the UMAP view."}), 400
    if cell_type or disease:
        return jsonify({"error": "Pre-generated UMAP images cannot be filtered by disease or cell type."}), 400
    image_path = get_static_umap_image_path(column)
    if image_path is None or not os.path.isfile(image_path):
        return jsonify({"error": "No pre-generated UMAP image is available for that metadata column."}), 404
    return jsonify({"column": column, "plot": url_for("umap_image", column=column)})


@app.route("/umap_image/<column>.png")
def umap_image(column: str):
    image_path = get_static_umap_image_path(column)
    if image_path is None or not os.path.isfile(image_path):
        return jsonify({"error": "Pre-generated UMAP image not found."}), 404
    return send_file(image_path, mimetype="image/png")


@app.route("/api/gene_expression")
def api_gene_expression():
    gene_query = request.args.get("gene", "").strip()
    cell_type = request.args.get("cell_type", "").strip()
    disease = request.args.get("disease", "").strip()
    plot_type = request.args.get("plot_type", "histogram").strip().lower()
    try:
        payload = get_gene_expression(gene_query, cell_type=cell_type or None, disease=disease or None, plot_type=plot_type)
        return jsonify(payload)
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 503
    except (KeyError, ValueError) as exc:
        return jsonify({"error": str(exc)}), 404


@app.route("/api/gene_violin")
def api_gene_violin():
    gene_query = request.args.get("gene", "").strip()
    cell_types = request.args.get("cell_types", "").strip()
    disease = request.args.get("disease", "").strip()
    try:
        payload = get_gene_violin(gene_query, cell_types=cell_types or None, disease=disease or None)
        return jsonify(payload)
    except FileNotFoundError as exc:
        return jsonify({"error": str(exc)}), 503
    except (KeyError, ValueError) as exc:
        return jsonify({"error": str(exc)}), 404


@app.route("/api/deg_summary")
def api_deg_summary():
    return jsonify(get_deg_summary())


@app.route("/api/volcano_catalog")
def api_volcano_catalog():
    catalog = get_volcano_catalog()
    return jsonify({"files": len(catalog), "plots": catalog})


@app.route("/volcano_image/<path:disease>.png")
def volcano_image(disease: str):
    target = get_volcano_path(disease)
    if target is None:
        return jsonify({"error": "Volcano plot not found."}), 404
    return send_file(target, mimetype="image/png")


@app.route("/api/deg_table")
def api_deg_table():
    disease = request.args.get("disease", "").strip()
    cell_type = request.args.get("cell_type", "").strip()
    regulation = request.args.get("regulation", "").strip()
    limit = max(0, int(request.args.get("limit", 0) or 0))
    sort_by = request.args.get("sort_by", "").strip()
    sort_order = request.args.get("sort_order", "desc").strip()
    return jsonify(get_deg_table(disease=disease or None, cell_type=cell_type or None, regulation=regulation or None, limit=limit, sort_by=sort_by or None, sort_order=sort_order))


@app.route("/api/deg_download")
def api_deg_download():
    disease = request.args.get("disease", "").strip()
    cell_type = request.args.get("cell_type", "").strip()
    regulation = request.args.get("regulation", "").strip()
    sort_by = request.args.get("sort_by", "").strip()
    sort_order = request.args.get("sort_order", "desc").strip()
    payload = get_deg_table(disease=disease or None, cell_type=cell_type or None, regulation=regulation or None, limit=0, sort_by=sort_by or None, sort_order=sort_order)
    frame = pd.DataFrame(payload["rows"])
    buffer = io.BytesIO()
    frame.to_csv(buffer, index=False)
    buffer.seek(0)
    parts = ["DEG", disease or "all", "all_cell_types" if cell_type == DEG_ALL_CELL_TYPES else (cell_type or "all_cell_types")]
    if regulation:
        parts.append(regulation)
    filename = re.sub(r"[^A-Za-z0-9_.+-]", "_", "_".join(parts)) + ".csv"
    return send_file(buffer, mimetype="text/csv", as_attachment=True, download_name=filename)


@app.route("/api/graph_summary")
def api_graph_summary():
    with BACKGROUND_DATA_LOCK:
        cached = BACKGROUND_DATA["graph_summary"]
        if cached is not None:
            return jsonify(cached)
    return jsonify(get_graph_summary())


@app.route("/api/ego_catalog")
def api_ego_catalog():
    return jsonify({"files": len(get_ego_catalog()), "graphs": get_ego_catalog()})


@app.route("/ego_image/<disease>/<cell_type>.png")
def ego_image(disease: str, cell_type: str):
    target = get_ego_image_path(disease, cell_type)
    if target is None:
        return jsonify({"error": "EGO graph not found."}), 404
    return send_file(target, mimetype="image/png")


@app.route("/api/cell_interaction_catalog")
def api_cell_interaction_catalog():
    kind = request.args.get("kind", "heatmap").strip().lower()
    if kind not in CELL_INTERACTION_KIND_SUFFIXES:
        kind = "heatmap"
    catalog = get_cell_interaction_catalog(kind)
    return jsonify({"files": len(catalog), "heatmaps": catalog, "cell_types": get_cell_interaction_cell_types()})


@app.route("/api/cell_interaction_pair")
def api_cell_interaction_pair():
    cell_type_a = request.args.get("cell_type_a", "").strip()
    cell_type_b = request.args.get("cell_type_b", "").strip()
    try:
        return jsonify(get_cell_interaction_pair_counts(cell_type_a, cell_type_b))
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400


@app.route("/cell_interaction/<disease>.png")
def cell_interaction_image(disease: str):
    kind = request.args.get("kind", "heatmap").strip().lower()
    if kind not in CELL_INTERACTION_KIND_SUFFIXES:
        kind = "heatmap"
    target = get_cell_interaction_path(disease, kind=kind)
    if target is None:
        return jsonify({"error": "Cell interaction image not found."}), 404
    as_attachment = request.args.get("download", "").strip().lower() in {"1", "true", "yes"}
    return send_file(
        target,
        mimetype="image/png",
        as_attachment=as_attachment,
        download_name=os.path.basename(target),
    )


@app.route("/api/graph_data")
def api_graph_data():
    disease = request.args.get("disease", "").strip()
    cell_type = request.args.get("cell_type", "").strip()
    gene_list = request.args.get("gene_list", "").strip()
    limit = int(request.args.get("limit", 50))
    return jsonify(get_graph_payload(disease=disease or None, cell_type=cell_type or None, limit=limit, gene_list=gene_list or None))


@app.route("/api/graph_download")
def api_graph_download():
    filename = request.args.get("file", "").strip()
    if not filename:
        return jsonify({"error": "No graph file specified."}), 400

    safe_name = os.path.basename(filename)
    if safe_name != filename:
        return jsonify({"error": "Invalid file name."}), 400

    target = os.path.join(GRAPH_ROOT, safe_name)
    if not os.path.isfile(target):
        return jsonify({"error": "Graph file not found."}), 404

    return send_file(target, as_attachment=True, download_name=safe_name)


@app.route("/api/graph_network")
def api_graph_network():
    disease = request.args.get("disease", "").strip()
    cell_type = request.args.get("cell_type", "").strip()
    gene_list = request.args.get("gene_list", "").strip()
    layout = request.args.get("layout", "random").strip()
    limit = int(request.args.get("limit", 50))
    try:
        return jsonify(get_graph_network_payload(disease=disease or None, cell_type=cell_type or None, limit=limit, gene_list=gene_list or None, layout=layout))
    except (ImportError, ValueError) as exc:
        return jsonify({"error": str(exc)}), 400


@app.route("/api/query", methods=["POST"])
def api_query():
    payload = request.get_json(silent=True) or {}
    query = str(payload.get("query") or "")
    result, status_code = call_openai(query, history=payload.get("history"))
    if status_code != 200:
        return jsonify(result), status_code
    return jsonify(result)
